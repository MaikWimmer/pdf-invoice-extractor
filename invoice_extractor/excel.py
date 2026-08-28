"""Write extracted invoices into a two-sheet workbook.

Sheet 1 is one row per invoice, sheet 2 one row per line item, joined by the
invoice number. That split is what makes the result usable: totals for
accounting, line items for anyone who has to ask what was actually bought.
"""
from __future__ import annotations

from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .model import Invoice

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
EURO = '#,##0.00 "€"'


def _style(sheet, widths) -> None:
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"


def write_workbook(invoices: Sequence[Invoice], target: Path) -> Path:
    """Write ``invoices`` to ``target`` and return the path."""
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Invoices"
    summary.append(["Source file", "Invoice number", "Date", "Customer",
                    "Net", "VAT", "Gross"])
    for invoice in invoices:
        summary.append([invoice.source_file, invoice.number, invoice.date,
                        invoice.customer, invoice.net, invoice.vat,
                        invoice.gross])
    for row in range(2, len(invoices) + 2):
        summary[f"C{row}"].number_format = "DD.MM.YYYY"
        for column in ("E", "F", "G"):
            summary[f"{column}{row}"].number_format = EURO
    _style(summary, {1: 26, 2: 18, 3: 12, 4: 26, 5: 12, 6: 12, 7: 12})

    items = workbook.create_sheet("Line items")
    items.append(["Invoice number", "Customer", "Description", "Quantity",
                  "Unit price", "Total"])
    row = 2
    for invoice in invoices:
        for item in invoice.items:
            items.append([invoice.number, invoice.customer, item.description,
                          item.quantity, item.unit_price, item.total])
            items[f"E{row}"].number_format = EURO
            items[f"F{row}"].number_format = EURO
            row += 1
    _style(items, {1: 18, 2: 26, 3: 34, 4: 10, 5: 13, 6: 13})

    target = Path(target)
    workbook.save(target)
    return target
